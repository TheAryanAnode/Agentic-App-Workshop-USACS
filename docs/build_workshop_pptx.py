#!/usr/bin/env python3
"""Build the USACS workshop deck on the Slidesgo 'AI in Marketing XL' theme.

Reads the official .pptx so layouts, Orbitron / Darker Grotesque, and 3D art
stay intact. Workshop copy is filled into cloned slides; unused template slides
are removed. Keep the Thanks slide for Slidesgo attribution.
"""

from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "slide-assets"
OUT = ROOT / "USACS-Tech-Committee-Agentic-App-Workshop.pptx"
TEMPLATE_CANDIDATES = [
    Path("/Users/aryan/Downloads/The Use of AI in Marketing XL by Slidesgo.pptx"),
    ROOT / "The Use of AI in Marketing XL by Slidesgo.pptx",
]

BLUE = (0x04, 0x45, 0xFF)
PINK = (0xF2, 0xCE, 0xFF)
CYAN = (0x8D, 0xF1, 0xFF)
GRAY = (0xE8, 0xE8, 0xE8)
INK = (0x19, 0x19, 0x19)
WHITE = (255, 255, 255)

NSMAP = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def find_template() -> Path:
    for path in TEMPLATE_CANDIDATES:
        if path.exists():
            return path
    raise FileNotFoundError(
        "Slidesgo template not found. Place "
        "'The Use of AI in Marketing XL by Slidesgo.pptx' in Downloads or docs/."
    )


def font_path(bold=False, mono=False):
    if mono:
        for p in (
            "/System/Library/Fonts/Menlo.ttc",
            "/System/Library/Fonts/Monaco.ttf",
            "/Library/Fonts/Menlo.ttc",
        ):
            if Path(p).exists():
                return p
    names = (
        ["Arial Bold.ttf", "Arial Bold.ttf"]
        if bold
        else ["Arial.ttf", "Arial Unicode.ttf"]
    )
    for folder in (
        Path("/System/Library/Fonts/Supplemental"),
        Path("/Library/Fonts"),
        Path("/System/Library/Fonts"),
    ):
        for name in names:
            p = folder / name
            if p.exists():
                return str(p)
    return None


def load_font(size, bold=False, mono=False):
    path = font_path(bold=bold, mono=mono)
    if path:
        try:
            return ImageFont.truetype(path, size, index=0)
        except Exception:
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                pass
    return ImageFont.load_default()


def rounded_rect(draw, box, radius, fill, outline=None, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def save_png(img, name):
    ASSETS.mkdir(exist_ok=True)
    path = ASSETS / name
    img.save(path)
    return path


def make_chatbot_vs_agent():
    w, h = 1200, 1380
    img = Image.new("RGB", (w, h), GRAY)
    d = ImageDraw.Draw(img)
    title = load_font(42, bold=True)
    body = load_font(28)
    small = load_font(22)

    def card(y, heading, lines, accent):
        rounded_rect(d, (70, y, 1130, y + 560), 40, WHITE, outline=accent, width=6)
        d.text((110, y + 36), heading, font=title, fill=accent)
        yy = y + 130
        for line in lines:
            rounded_rect(d, (110, yy, 1090, yy + 78), 18, (244, 244, 246))
            d.text((140, yy + 22), line, font=body, fill=INK)
            yy += 100

    card(70, "CHATBOT", ["Reads the prompt", "Guesses the next words", "Sounds sure. Can still be wrong."], BLUE)
    card(700, "AGENT", ["Has a goal", "Calls your tools", "Uses the real result, then answers"], (90, 70, 210))
    d.text((70, 1305), "Same model. Extra loop: think, act, observe.", font=small, fill=(90, 90, 98))
    return save_png(img, "chatbot-vs-agent-xl.png")


def make_langgraph():
    w, h = 1920, 1080
    img = Image.new("RGB", (w, h), GRAY)
    d = ImageDraw.Draw(img)
    title = load_font(36, bold=True)
    label = load_font(28, bold=True)
    small = load_font(22)

    boxes = [
        (80, 430, 280, 620, "START"),
        (430, 430, 760, 620, "AGENT"),
        (980, 180, 1320, 370, "TOOLS?"),
        (980, 690, 1320, 880, "ANSWER"),
        (1480, 180, 1820, 370, "RUN CODE"),
        (1480, 690, 1820, 880, "END"),
    ]
    for x1, y1, x2, y2, text in boxes:
        fill = WHITE
        outline = BLUE
        if text in ("START", "END"):
            fill = PINK
            outline = PINK
        if text == "RUN CODE":
            fill = CYAN
            outline = BLUE
        rounded_rect(d, (x1, y1, x2, y2), 28, fill, outline=outline, width=5)
        bbox = d.textbbox((0, 0), text, font=label)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        d.text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2 - 4), text, font=label, fill=INK)

    def arrow(a, b):
        d.line([a, b], fill=BLUE, width=6)

    arrow((280, 525), (430, 525))
    arrow((760, 500), (980, 275))
    arrow((760, 550), (980, 785))
    arrow((1320, 275), (1480, 275))
    arrow((1650, 370), (1650, 690))
    arrow((1480, 785), (1320, 785))
    arrow((1650, 275), (870, 275))
    d.line([(1650, 275), (870, 275), (870, 430)], fill=BLUE, width=6)

    d.text((80, 80), "LANGGRAPH IS THIS LOOP", font=title, fill=INK)
    d.text((80, 140), "Need a tool? Yes: run your function, then think again. No: answer and stop.", font=small, fill=(70, 70, 80))
    return save_png(img, "langgraph-loop-xl.png")


def make_code_png(lines, name, width=1100, height=1600):
    img = Image.new("RGB", (width, height), INK)
    d = ImageDraw.Draw(img)
    mono = load_font(28, mono=True)
    y = 70
    for line in lines:
        color = WHITE
        stripped = line.strip()
        if stripped.startswith("#") or stripped.startswith("//"):
            color = CYAN
        elif stripped.startswith("import") or stripped.startswith("export") or stripped.startswith("const"):
            color = PINK
        d.text((70, y), line if line else " ", font=mono, fill=color)
        y += 46
    return save_png(img, name)


def remap_rids(element, mapping):
    for el in element.iter():
        for attr in list(el.attrib):
            if attr.endswith("}embed") or attr.endswith("}link"):
                old = el.get(attr)
                if old in mapping:
                    el.set(attr, mapping[old])


def copy_rel(reltype):
    name = reltype.rsplit("/", 1)[-1]
    return name in {"image", "hyperlink", "chart", "diagram", "oleObject", "video", "media", "audio"}


def duplicate_slide(prs, index):
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    sp_tree = dest.shapes._spTree
    for child in list(sp_tree):
        tag = child.tag
        if tag.endswith("}nvGrpSpPr") or tag.endswith("}grpSpPr"):
            continue
        sp_tree.remove(child)

    rid_map = {}
    for r_id, rel in source.part.rels.items():
        if not copy_rel(rel.reltype):
            continue
        if rel.is_external:
            new_rid = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = dest.part.rels.get_or_add(rel.reltype, rel.target_part)
        rid_map[r_id] = new_rid

    for child in source.shapes._spTree:
        tag = child.tag
        if tag.endswith("}nvGrpSpPr") or tag.endswith("}grpSpPr"):
            continue
        newel = deepcopy(child)
        remap_rids(newel, rid_map)
        sp_tree.append(newel)

    src_csld = source._element.find(qn("p:cSld"))
    dst_csld = dest._element.find(qn("p:cSld"))
    src_bg = src_csld.find(qn("p:bg"))
    if src_bg is not None:
        old = dst_csld.find(qn("p:bg"))
        if old is not None:
            dst_csld.remove(old)
        new_bg = deepcopy(src_bg)
        remap_rids(new_bg, rid_map)
        dst_csld.insert(0, new_bg)

    return dest


def delete_slide(prs, index):
    sld_id_lst = prs.slides._sldIdLst
    sld_id = sld_id_lst[index]
    r_id = sld_id.get(qn("r:id"))
    prs.part.drop_rel(r_id)
    sld_id_lst.remove(sld_id)


def set_paragraphs(shape, lines):
    tf = shape.text_frame
    tf.word_wrap = True
    while len(tf.paragraphs) < len(lines):
        tf.add_paragraph()
    paras = list(tf.paragraphs)
    for i, line in enumerate(lines):
        p = paras[i]
        if p.runs:
            p.runs[0].text = line
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.text = line
    for p in paras[len(lines) :]:
        if p.runs:
            p.runs[0].text = ""
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.text = ""


def set_text(shape, text):
    set_paragraphs(shape, [text] if isinstance(text, str) else text)


def ph(slide, idx):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    raise KeyError(idx)


def shapes_with_text(slide, needle):
    found = []
    for shape in slide.shapes:
        if shape.has_text_frame and needle in shape.text_frame.text:
            found.append(shape)
    return found


def replace_picture(shape, path):
    image_part, r_id = shape.part.get_or_add_image_part(str(path))
    blip = shape._element.find(".//" + qn("a:blip"))
    if blip is None:
        raise RuntimeError("shape has no blip")
    blip.set(qn("r:embed"), r_id)


def largest_picture(slide):
    pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not pics:
        return None
    return max(pics, key=lambda s: int(s.width) * int(s.height))


def replace_background(slide, path):
    image_part, r_id = slide.part.get_or_add_image_part(str(path))
    bg = slide._element.find(qn("p:cSld")).find(qn("p:bg"))
    if bg is None:
        return
    blip = bg.find(".//" + qn("a:blip"))
    if blip is not None:
        blip.set(qn("r:embed"), r_id)


def set_page_marker(slide, n):
    """Update the decorative droplet number, not numbered TOC/section placeholders."""
    for shape in slide.shapes:
        if not shape.has_text_frame or shape.is_placeholder:
            continue
        text = shape.text_frame.text.strip()
        if text.isdigit() and len(text) <= 2:
            set_text(shape, str(n))
            return


def notes(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    tf.clear()
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(14)
    run.font.name = "Calibri"


def fill_pair(slide, original_a, original_b, line1, line2):
    for shape in shapes_with_text(slide, original_a):
        set_text(shape, line1)
    for shape in shapes_with_text(slide, original_b):
        set_text(shape, line2)


def build():
    template = find_template()
    make_chatbot_vs_agent()
    make_langgraph()
    graph_code = make_code_png(
        [
            "const workflow = new StateGraph(AgentState)",
            '  .addNode("agent", agentNode)',
            '  .addNode("tools", toolNode)',
            '  .addEdge(START, "agent")',
            '  .addConditionalEdges(',
            '    "agent", toolsCondition,',
            '    ["tools", END]',
            "  )",
            '  .addEdge("tools", "agent")',
        ],
        "code-graph.png",
    )
    dining_code = make_code_png(
        [
            'import { searchDining } from "@/data/dining"',
            "",
            "export const findDining = tool(",
            "  async ({ query }) => {",
            "    const matches = searchDining(query)",
            '    return JSON.stringify({',
            '      kind: "dining", query, matches',
            "    })",
            "  },",
            "  { name: \"findDining\" }",
            ")",
            "export const tools = [",
            "  getCourseInfo, findDining",
            "]",
        ],
        "code-dining.png",
    )

    prs = Presentation(str(template))
    orig_count = len(prs.slides)

    def take(src):
        return duplicate_slide(prs, src)

    slides = []

    def keep(slide, note, number=None):
        notes(slide, note)
        if number is not None:
            set_page_marker(slide, number)
        slides.append(slide)
        return slide

    # 1 Title
    s = take(0)
    fill_pair(s, "THE USE OF AI", "IN MARKETING", "AGENTIC APP", "WORKSHOP")
    set_text(ph(s, 0), "AGENTIC APP WORKSHOP")
    set_text(ph(s, 1), "USACS Tech Committee  ·  Workshop 1")
    keep(
        s,
        "Welcome people in. This is USACS Tech Committee, workshop 1. They leave with a working campus agent. Theme is the Slidesgo AI in Marketing XL template (Orbitron titles, Darker Grotesque body).",
        1,
    )

    # 2 WHOA / you build
    s = take(2)
    for shape in list(s.shapes):
        if not shape.has_text_frame or "WHOA" not in shape.text_frame.text:
            continue
        set_text(shape, "YOU BUILD" if shape.is_placeholder else "BUILD")
    set_paragraphs(
        ph(s, 1),
        [
            "You add tools to a real app.",
            "You watch the agent call your code.",
            "You do not sit through a lecture.",
        ],
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "students.jpg")
    keep(
        s,
        "Set the contract. Pair people with slow laptops. Ask who has never used an API key and tell them that is fine.",
        2,
    )

    # 3 TOC
    s = take(4)
    set_text(ph(s, 21), "TABLE OF CONTENTS")
    toc = [
        (2, "01", 0, "PROBLEM", 1, "A real Rutgers question"),
        (4, "02", 3, "THE LOOP", 5, "Think, act, observe"),
        (7, "03", 6, "TOOLS", 8, "Functions the model can call"),
        (13, "04", 9, "BUILD", 14, "Dining, buildings, GPA"),
        (16, "05", 15, "FINAL BOSS", 17, "Course + food on Busch"),
        (19, "06", 18, "SHIP IT", 20, "What you walk out with"),
    ]
    for num_idx, num, title_idx, title, body_idx, body in toc:
        set_text(ph(s, num_idx), num)
        set_text(ph(s, title_idx), title)
        set_text(ph(s, body_idx), body)
    keep(s, "Two minutes. Tonight is not a theory talk. They will run an app, then write tools.", 3)

    # 4 Section 01
    s = take(8)
    set_paragraphs(ph(s, 0), ["THE", "PROBLEM"])
    set_text(ph(s, 1), "Start with something they actually want answered.")
    set_text(ph(s, 2), "01")
    keep(s, "Pause. Then read the prompt on the next slide out loud.", 4)

    # 5 Prompt + vote
    s = take(9)
    set_text(ph(s, 0), "A REAL RUTGERS QUESTION")
    set_paragraphs(
        ph(s, 1),
        [
            "CS112 tomorrow. Food on Busch after.",
            "Find the course. GPA for A, B+, and B. Dining nearby.",
            "",
            "Raise a hand. Could a normal chatbot actually do this?",
            "Yes: it can sound confident.",
            "No: it cannot look anything up unless you give it tools.",
        ],
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "campus.jpg")
    keep(s, "Read the prompt out loud. Take a real vote. This prompt is the north star for the Final Boss.", 5)

    # 6 Why agent
    s = take(16)
    set_text(ph(s, 0), "WHY AN AGENT, NOT A CHATBOT")
    set_text(
        ph(s, 1),
        "A chatbot predicts words. An agent has a goal, calls your code, and keeps going.",
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "chatbot-vs-agent-xl.png")
    keep(
        s,
        "Walk the picture. Chatbot path is guess text. Agent path is call a tool, then use the real result. That is the whole why.",
        6,
    )

    # 7 Think / act
    s = take(20)
    fill_pair(s, "AWESOME", "WORDS", "THINK", "ACT")
    set_text(ph(s, 0), "THINK ACT")
    keep(
        s,
        "Think. Act. Observe. Repeat until you can answer. Same loop as Cursor, support bots, and research agents. Tonight it is Rutgers.",
        7,
    )

    # 8 Tools (3 columns)
    s = take(14)
    set_text(ph(s, 6), "TOOLS ARE FUNCTIONS IT CAN CALL")
    set_text(ph(s, 0), "CS112")
    set_text(ph(s, 1), "getCourseInfo")
    set_text(ph(s, 2), "BUSCH FOOD")
    set_text(ph(s, 3), "findDining")
    set_text(ph(s, 4), "GPA")
    set_text(ph(s, 5), "calculateGrade")
    keep(
        s,
        "The model picks a name and arguments. Your TypeScript runs. If it is not in the tools array, the agent cannot call it. Building and events work the same way.",
        8,
    )

    # 9 LangGraph full-bleed
    s = take(18)
    set_text(ph(s, 0), "LANGGRAPH IS THIS LOOP")
    replace_background(s, ASSETS / "langgraph-loop-xl.png")
    keep(
        s,
        "Do not teach the LangGraph API. Teach the picture. START to agent. Need a tool? Yes goes to run code and back. No goes to answer.",
        9,
    )

    # 10 graph.ts
    s = take(10)
    set_text(ph(s, 0), "THE SAME LOOP IN GRAPH.TS")
    set_paragraphs(
        ph(s, 1),
        [
            "Read it. You do not have to write it.",
            "START goes to agent.",
            "toolsCondition is yes or no on tool_calls.",
            "tools then back to agent.",
            "That is tonight’s LangGraph.",
        ],
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, graph_code)
    keep(s, "Walk line by line. This is the same picture they just saw, in code.", 10)

    # 11 App is built
    s = take(17)
    set_text(ph(s, 0), "THE APP IS BUILT. YOU ADD SKILLS.")
    set_text(ph(s, 1), "Live in agent/tools.ts. data/ has courses, dining, buildings, events. Ignore the rest until a tool works.")
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "laptop.jpg")
    keep(
        s,
        "Anxiety reducer. They will not write Next.js streaming. Click tools.ts. Point at export const tools = [getCourseInfo].",
        11,
    )

    # 12 Start
    s = take(29)
    set_text(ph(s, 0), "START THE APP")
    set_text(ph(s, 1), "cp .env.example .env.local   then  ./run.sh")
    set_text(ph(s, 2), "02")
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "laptop.jpg")
    keep(
        s,
        "Walk the room. Google AI Studio keys. Do not share one key for 50 people. Bypass starts the UI but chat will fail until a key exists.",
        12,
    )

    # 13 Challenge 1
    s = take(9)
    set_text(ph(s, 0), "CHALLENGE 1  WATCH A TOOL FIRE")
    set_paragraphs(
        ph(s, 1),
        [
            "Ask: What is CS112 and where is it usually taught?",
            "",
            "Want getCourseInfo in the log and a course card in chat.",
            "This tool is already done.",
            "Copy later: body, name, description, schema.",
            "JSON.stringify keeps Gemini happy.",
        ],
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "collaboration.jpg")
    keep(s, "Give 3 minutes. If a tool never appears, it is not registered.", 13)

    # 14 Activity log
    s = take(44)
    set_text(ph(s, 8), "THE RIGHT PANEL IS LANGGRAPH")
    set_text(ph(s, 0), "REQUEST")
    set_text(ph(s, 1), "Request received")
    set_text(ph(s, 2), "DECIDE")
    set_text(ph(s, 3), "Agent deciding")
    set_text(ph(s, 4), "CALL")
    set_text(ph(s, 5), "Calling getCourseInfo")
    set_text(ph(s, 6), "ANSWER")
    set_text(ph(s, 7), "Response generated")
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "students.jpg")
    keep(
        s,
        "If a tool never appears, it is not registered. If it errors, their function threw or returned bad JSON. Debug the log first.",
        14,
    )

    # 15 Dining
    s = take(16)
    set_text(ph(s, 0), "CHALLENGE 2  GIVE IT DINING")
    set_text(
        ph(s, 1),
        "Ask: Where can I eat on Busch? Import searchDining. Return kind: dining. Add findDining to the tools array.",
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "campus.jpg")
    keep(
        s,
        "This is the aha. Same app, missing skill. Give 10 minutes. Two bugs: forgot the import, forgot the array. Do not show the answer key until most people have tried.",
        15,
    )

    # 16 Answer key
    s = take(10)
    set_text(ph(s, 0), "ANSWER KEY  FIND DINING")
    set_paragraphs(
        ph(s, 1),
        [
            "Reveal after they attempt it.",
            "searchDining already filters.",
            "The last line is half of the bug reports.",
            "findBuilding is the same shape.",
        ],
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, dining_code)
    keep(s, "Leave this up while they catch up. The last line is half of the bug reports.", 16)

    # 17 Further
    s = take(44)
    set_text(ph(s, 8), "THEN GO FURTHER")
    set_text(ph(s, 0), "BUILDING")
    set_text(ph(s, 1), "Nicknames like Hill and CORE.")
    set_text(ph(s, 2), "EVENTS")
    set_text(ph(s, 3), "Filter by CS or campus.")
    set_text(ph(s, 4), "GRADE")
    set_text(ph(s, 5), "Parse A, B+, B into a GPA.")
    set_text(ph(s, 6), "BREAK IT")
    set_text(ph(s, 7), "Leave it out of the tools array.")
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "collaboration.jpg")
    keep(
        s,
        "Same pattern as dining. Fast people do building and events. Grade is needed for the Final Boss. Normalize failure: I implemented it and forgot the array.",
        17,
    )

    # 18 Final boss
    s = take(16)
    set_text(ph(s, 0), "FINAL BOSS")
    set_text(
        ph(s, 1),
        "CS112 tomorrow. Food on Busch after. Course info. GPA for A, B+, B. Dining nearby. A complete agent calls all three.",
    )
    pic = largest_picture(s)
    if pic:
        replace_picture(pic, ASSETS / "collaboration.jpg")
    keep(
        s,
        "Bring the room back. Run this on a volunteer laptop. Narrate the log. If a tool is missing, the agent uses the ones it has.",
        18,
    )

    # 19 What you built
    s = take(45)
    set_text(ph(s, 8), "WHAT YOU BUILT")
    set_text(ph(s, 0), "LLM")
    set_text(ph(s, 1), "Decides which action to take.")
    set_text(ph(s, 2), "TOOLS")
    set_text(ph(s, 3), "Give the agent real capabilities.")
    set_text(ph(s, 4), "LANGGRAPH")
    set_text(ph(s, 5), "Runs the request, tool, result loop.")
    set_text(ph(s, 6), "UI")
    set_text(ph(s, 7), "Turns JSON into cards you can see.")
    keep(
        s,
        "Have the room say it back. You do not need to know everything. You built something real. You can build the next thing. That is the USACS Tech Committee bar.",
        19,
    )

    # 20 Thanks + required attribution
    s = take(59)
    set_text(ph(s, 0), "THANKS")
    for shape in shapes_with_text(s, "THANKS"):
        set_text(shape, "THANKS")
    set_paragraphs(
        ph(s, 1),
        [
            "Stay for help. Keep building.",
            "",
            "github.com/TheAryanAnode/",
            "Agentic-App-Workshop-USACS",
            "USACS Tech Committee",
        ],
    )
    keep(
        s,
        "Leave this up during Q and A. Solution branch has completed tools. Remind them not to commit .env.local. Keep this slide: Slidesgo free license requires attribution.",
        20,
    )

    for _ in range(orig_count):
        delete_slide(prs, 0)

    if len(prs.slides) != 20:
        raise SystemExit(f"expected 20 slides, got {len(prs.slides)}")

    prs.save(str(OUT))
    print(len(prs.slides), OUT)


if __name__ == "__main__":
    build()
